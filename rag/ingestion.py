from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path
from typing import BinaryIO

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from .config import Settings

CATEGORIES = ("hr", "technical", "compliance")
HEADING_RE = re.compile(r"(?m)^\s*(?:\d+\.\s+)?([A-Z][A-Z0-9 /&()\-]{3,})\s*$")


def infer_category(path: Path) -> str:
    lowered = " ".join(part.lower() for part in path.parts)
    for category in CATEGORIES:
        if category in lowered:
            return category
    return "general"


def extract_section(text: str) -> str:
    matches = HEADING_RE.findall(text)
    return matches[0].strip().title() if matches else "General"


def read_path(path: Path) -> list[tuple[str, int | None]]:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return [(path.read_text(encoding="utf-8", errors="replace"), None)]
    if suffix == ".pdf":
        return [(page.extract_text() or "", number) for number, page in enumerate(PdfReader(str(path)).pages, 1)]
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        return [("\n".join(p.text for p in doc.paragraphs), None)]
    raise ValueError(f"Unsupported file type: {suffix}")


def read_upload(name: str, stream: BinaryIO) -> list[tuple[str, int | None]]:
    suffix = Path(name).suffix.lower()
    payload = stream.read()
    if suffix == ".txt":
        return [(payload.decode("utf-8", errors="replace"), None)]
    if suffix == ".pdf":
        return [(page.extract_text() or "", number) for number, page in enumerate(PdfReader(io.BytesIO(payload)).pages, 1)]
    if suffix == ".docx":
        doc = DocxDocument(io.BytesIO(payload))
        return [("\n".join(p.text for p in doc.paragraphs), None)]
    raise ValueError(f"Unsupported file type: {suffix}")


def make_documents(parts: list[tuple[str, int | None]], source: str, category: str, cfg: Settings) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=cfg.chunk_size, chunk_overlap=cfg.chunk_overlap)
    documents: list[Document] = []
    for text, page in parts:
        for index, chunk in enumerate(splitter.split_text(text)):
            if not chunk.strip():
                continue
            digest = hashlib.sha256(f"{source}:{page}:{index}:{chunk}".encode()).hexdigest()[:20]
            documents.append(Document(page_content=chunk, metadata={
                "id": digest, "source": source, "category": category,
                "section": extract_section(chunk), "page": page or 0,
            }))
    return documents


def bundled_documents(cfg: Settings) -> list[Document]:
    documents: list[Document] = []
    for path in sorted(cfg.docs_dir.rglob("*")):
        if path.suffix.lower() in {".txt", ".pdf", ".docx"}:
            documents.extend(make_documents(read_path(path), path.name, infer_category(path), cfg))
    return documents

