from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Enterprise_RAG_Bot_Architecture.docx"
DIAGRAM = ROOT / "architecture_diagram.png"

NAVY = "123653"
TEAL = "167D9A"
LIGHT = "EAF2F7"
MID = "D4E3EC"
INK = "183042"
MUTED = "5E7180"
WHITE = "FFFFFF"


def font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, detail, fill, outline=TEAL):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline="#" + outline, width=3)
    x1, y1, x2, y2 = xy
    draw.text(((x1+x2)//2, y1+24), title, font=font(26, True), fill="#" + INK, anchor="ma")
    lines = detail.split("\n")
    for i, line in enumerate(lines):
        draw.text(((x1+x2)//2, y1+70+i*28), line, font=font(19), fill="#" + MUTED, anchor="ma")


def arrow(draw, start, end, label=None):
    draw.line([start, end], fill="#" + TEAL, width=5)
    x2, y2 = end
    x1, y1 = start
    import math
    angle = math.atan2(y2-y1, x2-x1)
    for delta in (2.55, -2.55):
        p = (x2 + 18*math.cos(angle+delta), y2 + 18*math.sin(angle+delta))
        draw.line([end, p], fill="#" + TEAL, width=5)
    if label:
        mx, my = (x1+x2)//2, (y1+y2)//2
        draw.text((mx, my-10), label, font=font(16, True), fill="#" + TEAL, anchor="ms")


def build_diagram():
    image = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(image)
    d.text((800, 42), "Enterprise RAG Bot - Logical Architecture", font=font(34, True), fill="#" + NAVY, anchor="ma")
    rounded_box(d, (80, 150, 360, 300), "Users", "Employees | Engineers\nRisk | Knowledge Admins", "#F4F8FA")
    rounded_box(d, (500, 125, 820, 325), "Streamlit UI", "Ingestion workspace\nQuestion answering\nSession history + evidence", "#E6F2F5")
    rounded_box(d, (970, 125, 1310, 325), "LangGraph", "Route -> Retrieve -> Grade\nGenerate or Refuse", "#EAF2F7")
    arrow(d, (360, 225), (500, 225), "HTTPS / local")
    arrow(d, (820, 225), (970, 225), "workflow state")

    rounded_box(d, (120, 500, 440, 700), "Document Pipeline", "TXT | PDF | DOCX loaders\nSection metadata\n900-char chunks / 150 overlap", "#F5F7F9")
    rounded_box(d, (610, 500, 930, 700), "Knowledge Base", "Provider embeddings\nChroma vector persistence\nTop-k=5 + category filter", "#E6F2F5")
    rounded_box(d, (1100, 500, 1450, 700), "Model Provider", "Ollama (local default)\nllama3.2:3b + nomic\nOpenAI remains supported", "#F5F7F9")
    arrow(d, (440, 600), (610, 600), "documents")
    arrow(d, (930, 600), (1100, 600), "embed / infer")
    arrow(d, (1140, 325), (1140, 500), "LLM calls")
    arrow(d, (970, 290), (850, 500), "retrieval")

    d.rounded_rectangle((80, 790, 1520, 850), radius=15, fill="#" + NAVY)
    d.text((800, 820), "Trust boundary: environment configuration + local model service + persistent index",
           font=font(21, True), fill="white", anchor="mm")
    image.save(DIAGRAM, quality=95)


def set_cell_fill(cell, color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, widths):
    table.style = "Table Grid"
    table_geometry(table, widths)
    for i, cell in enumerate(table.rows[0].cells):
        set_cell_fill(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9.5)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def build_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.85)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.header_distance = sec.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in [
        ("Heading 1", 16, NAVY, 16, 8), ("Heading 2", 13, TEAL, 12, 6), ("Heading 3", 11.5, NAVY, 8, 4)
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = sec.header.paragraphs[0]
    header.text = "ACME KNOWLEDGE SYSTEMS  |  ARCHITECTURE"
    header.style = styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ENTERPRISE RAG BOT")
    r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Architecture Document")
    r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph("Local-first retrieval-augmented question answering for HR, technical, and compliance knowledge")
    p.paragraph_format.space_after = Pt(18)
    p.runs[0].font.size = Pt(13); p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    meta = doc.add_table(rows=4, cols=2)
    data = [("System", "Enterprise RAG Bot"), ("Architecture status", "Current implementation"),
            ("Primary runtime", "Streamlit + LangGraph + Chroma + Ollama"), ("Document date", date.today().isoformat())]
    for row, values in zip(meta.rows, data):
        row.cells[0].text, row.cells[1].text = values
        set_cell_fill(row.cells[0], LIGHT)
        row.cells[0].paragraphs[0].runs[0].bold = True
    table_geometry(meta, [2700, 6660])

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Purpose")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph("Define the implemented system boundaries, components, data flows, deployment assumptions, trust controls, and extension points for engineering, operations, and security reviewers.")
    doc.add_page_break()

    add_heading(doc, "1. Executive Summary")
    doc.add_paragraph("The Enterprise RAG Bot is a local-first knowledge assistant that answers questions from approved enterprise documents. It separates corpus administration from end-user question answering and exposes source evidence and workflow status for each answer.")
    add_bullet(doc, "Three knowledge domains: HR, Technical, and Compliance.")
    add_bullet(doc, "Two provider modes: Ollama for local inference and OpenAI as a supported alternative.")
    add_bullet(doc, "Persistent semantic retrieval through a Chroma collection with category-aware filtering.")
    add_bullet(doc, "A LangGraph state machine controls routing, retrieval, evidence grading, answer generation, and safe refusal.")
    add_bullet(doc, "Streamlit owns presentation and session history; model and retrieval concerns remain behind Python boundaries.")

    add_heading(doc, "2. Architecture Principles")
    principles = [
        ("Ground answers in evidence", "The model receives numbered excerpts and is instructed to cite them."),
        ("Fail safely", "Weak or missing evidence routes to a refusal instead of an unsupported answer."),
        ("Keep providers replaceable", "Configuration selects Ollama or OpenAI without changing the UI workflow."),
        ("Preserve provenance", "Source, category, section, page, and stable chunk ID travel with every chunk."),
        ("Separate local indexes", "Ollama and OpenAI indexes use distinct persistence paths to avoid embedding-dimension conflicts."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Principle", "Implementation"
    for a, b in principles:
        cells = table.add_row().cells; cells[0].text = a; cells[1].text = b
    style_table(table, [2700, 6660])
    doc.add_page_break()

    add_heading(doc, "3. Logical Architecture")
    doc.add_paragraph("The diagram shows the major runtime components and the two principal flows: document ingestion and question answering.")
    doc.add_picture(str(DIAGRAM), width=Inches(6.45))
    cap = doc.add_paragraph("Figure 1. Logical architecture and trust boundary")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True; cap.runs[0].font.size = Pt(9); cap.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    add_heading(doc, "4. Component Responsibilities")
    rows = [
        ("Streamlit UI", "app.py", "Tabs, uploads, controls, chat history, evidence display, and workflow trace."),
        ("Configuration", "rag/config.py", "Loads environment settings, provider/model selection, index path, chunking, and retrieval values."),
        ("Ingestion", "rag/ingestion.py", "Loads TXT/PDF/DOCX, extracts sections, chunks text, and creates stable metadata-rich documents."),
        ("Knowledge base", "rag/store.py", "Selects embeddings, persists Chroma data, adds/searches chunks, counts records, and resets the index."),
        ("Workflow", "rag/workflow.py", "Implements route, retrieve, grade, generate, and refuse nodes in LangGraph."),
        ("Model runtime", "Ollama/OpenAI", "Creates embeddings and performs evidence grading and grounded answer generation."),
    ]
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Component", "Location", "Responsibility")): cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row): cell.text = text
    style_table(table, [1900, 2100, 5360])
    doc.add_page_break()

    add_heading(doc, "5. Runtime Workflows")
    add_heading(doc, "5.1 Document ingestion", 2)
    for text in [
        "An administrator selects bundled documents or uploads approved TXT, PDF, or DOCX files.",
        "The loader extracts text and page information; the category is inferred from the bundled path or selected in the UI.",
        "The splitter creates 900-character chunks with 150-character overlap.",
        "Each chunk receives a stable SHA-256-derived ID plus source, category, section, and page metadata.",
        "The configured embedding provider creates vectors and Chroma persists them in the provider-specific index directory.",
    ]: add_bullet(doc, text)

    add_heading(doc, "5.2 Question answering", 2)
    for text in [
        "The user submits a question and may select a knowledge domain.",
        "The route node honors an explicit domain or applies keyword-based category routing.",
        "The retrieve node performs top-k semantic search (k=5), applying a category filter when present.",
        "The grade node asks the configured chat model whether the evidence can answer the question.",
        "Sufficient evidence proceeds to grounded generation with numbered citations; insufficient evidence produces a safe refusal.",
        "The UI renders the answer, retrieved chunks, provenance, and graph trace, then stores the turn in session state.",
    ]: add_bullet(doc, text)

    add_heading(doc, "6. Data Model and Persistence")
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Field", "Example", "Purpose")): cell.text = text
    for row in [
        ("id", "20-char content digest", "Stable update identity"), ("source", "hr_policy.txt", "Citation provenance"),
        ("category", "hr", "Retrieval filtering"), ("section", "Annual Leave", "Human-readable evidence context"),
        ("page", "3 or 0", "Page citation where available"), ("page_content", "Chunk text", "Embedding and grounded context"),
    ]:
        cells = table.add_row().cells
        for cell, text in zip(cells, row): cell.text = text
    style_table(table, [1500, 2700, 5160])

    doc.add_page_break()
    add_heading(doc, "7. Deployment and Configuration")
    doc.add_paragraph("The reference deployment runs as a local Streamlit process on port 8501 and connects to Ollama on localhost:11434. Ollama is the configured local provider; OpenAI remains available through environment configuration.")
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Variable", "Current/default", "Effect")): cell.text = text
    for row in [
        ("MODEL_PROVIDER", "ollama", "Selects Ollama or OpenAI adapters"),
        ("OLLAMA_CHAT_MODEL", "llama3.2:3b", "Evidence grading and answers"),
        ("OLLAMA_EMBEDDING_MODEL", "nomic-embed-text", "Document/query vectors"),
        ("OLLAMA_BASE_URL", "http://localhost:11434", "Local Ollama endpoint"),
        ("RAG_INDEX_DIR", ".rag_index_ollama", "Optional persistence override"),
        ("OPENAI_API_KEY", "Required only for OpenAI", "Cloud-provider authentication"),
    ]:
        cells = table.add_row().cells
        for cell, text in zip(cells, row): cell.text = text
    style_table(table, [2500, 2700, 4160])

    add_heading(doc, "8. Security and Trust Boundaries")
    add_bullet(doc, "Retrieved document text is treated as untrusted data, not executable instruction.")
    add_bullet(doc, "Secrets remain in environment configuration and are not rendered in the UI or document metadata.")
    add_bullet(doc, "The local Ollama mode keeps inference on the workstation, subject to local endpoint and file-system controls.")
    add_bullet(doc, "The current reference build does not implement authentication, document-level authorization, tenant isolation, encryption policy, malware scanning, or audit logging.")
    add_bullet(doc, "Production retrieval filters must enforce user/document access at query time; UI-only controls are insufficient.")

    add_heading(doc, "9. Reliability, Testing, and Operations")
    doc.add_paragraph("The repository currently includes fast unit coverage for category inference, section extraction, and heuristic routing. Operational readiness should add end-to-end ingestion/query tests, retrieval recall evaluations, faithfulness scoring, model availability checks, structured logging, backups, and index rebuild procedures.")

    add_heading(doc, "10. Recommended Production Evolution")
    for text in [
        "Add SSO/RBAC and document-level access metadata enforced by the retrieval layer.",
        "Move Chroma to an encrypted managed vector service with tenant isolation and backups.",
        "Add ingestion controls: malware scanning, DLP/PII classification, retention, and approval workflow.",
        "Introduce tracing, latency/error metrics, prompt/model versioning, rate limits, and user feedback.",
        "Maintain curated evaluation sets for retrieval recall, answer faithfulness, refusal quality, and security regressions.",
    ]: add_bullet(doc, text)

    add_heading(doc, "Appendix A. Repository Map")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Path", "Purpose"
    for row in [("app.py", "Streamlit application"), ("rag/config.py", "Environment-backed settings"),
                ("rag/ingestion.py", "Loaders and chunking"), ("rag/store.py", "Embedding and Chroma boundary"),
                ("rag/workflow.py", "LangGraph orchestration"), ("docs/", "Bundled enterprise corpus"),
                ("tests/", "Unit tests"), (".rag_index_ollama/", "Local Ollama vector index")]:
        cells = table.add_row().cells; cells[0].text, cells[1].text = row
    style_table(table, [2800, 6560])

    core = doc.core_properties
    core.title = "Enterprise RAG Bot Architecture"
    core.subject = "Logical, data, runtime, deployment, and security architecture"
    core.author = "ACME Knowledge Systems"
    doc.save(OUT)


if __name__ == "__main__":
    build_diagram()
    build_doc()
    print(OUT)
